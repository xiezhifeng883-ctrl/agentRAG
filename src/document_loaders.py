"""
Document loaders for various file formats.
Supports PDF, DOCX, PPTX, XLSX, TXT, and MD files.
"""

import os
from pathlib import Path
from typing import List, Dict, Optional
from abc import ABC, abstractmethod

import PyPDF2
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook
import markdown


class DocumentLoader(ABC):
    """Abstract base class for document loaders."""

    @abstractmethod
    def load(self, file_path: str) -> Dict[str, any]:
        """
        Load a document and return its content.

        Args:
            file_path: Path to the document file

        Returns:
            Dictionary with 'text', 'metadata', and 'file_path'
        """
        pass

    @staticmethod
    def get_loader(file_path: str) -> 'DocumentLoader':
        """
        Get the appropriate loader based on file extension.

        Args:
            file_path: Path to the document file

        Returns:
            DocumentLoader instance

        Raises:
            ValueError: If file format is not supported
        """
        extension = Path(file_path).suffix.lower()

        loaders = {
            '.pdf': PDFLoader,
            '.docx': DOCXLoader,
            '.doc': DOCXLoader,
            '.pptx': PPTXLoader,
            '.ppt': PPTXLoader,
            '.xlsx': XLSXLoader,
            '.xls': XLSXLoader,
            '.txt': TXTLoader,
            '.md': MarkdownLoader,
        }

        loader_class = loaders.get(extension)
        if not loader_class:
            raise ValueError(f"Unsupported file format: {extension}")

        return loader_class()


class PDFLoader(DocumentLoader):
    """Loader for PDF documents."""

    def load(self, file_path: str) -> Dict[str, any]:
        """Load PDF document."""
        text_content = []
        metadata = {
            'source': file_path,
            'file_type': 'pdf',
            'num_pages': 0
        }

        try:
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                metadata['num_pages'] = len(pdf_reader.pages)

                for page_num, page in enumerate(pdf_reader.pages):
                    text = page.extract_text()
                    if text.strip():
                        text_content.append(f"[Page {page_num + 1}]\n{text}")

                # Extract PDF metadata if available
                if pdf_reader.metadata:
                    metadata['title'] = pdf_reader.metadata.get('/Title', '')
                    metadata['author'] = pdf_reader.metadata.get('/Author', '')
                    metadata['subject'] = pdf_reader.metadata.get('/Subject', '')

        except Exception as e:
            raise RuntimeError(f"Error loading PDF {file_path}: {str(e)}")

        return {
            'text': '\n\n'.join(text_content),
            'metadata': metadata,
            'file_path': file_path
        }


class DOCXLoader(DocumentLoader):
    """Loader for DOCX documents."""

    def load(self, file_path: str) -> Dict[str, any]:
        """Load DOCX document."""
        text_content = []
        metadata = {
            'source': file_path,
            'file_type': 'docx'
        }

        try:
            doc = Document(file_path)

            # Extract core properties
            core_props = doc.core_properties
            metadata['title'] = core_props.title or ''
            metadata['author'] = core_props.author or ''
            metadata['subject'] = core_props.subject or ''

            # Extract paragraphs
            for para in doc.paragraphs:
                if para.text.strip():
                    text_content.append(para.text)

            # Extract tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = ' | '.join(cell.text.strip() for cell in row.cells)
                    if row_text.strip():
                        text_content.append(row_text)

        except Exception as e:
            raise RuntimeError(f"Error loading DOCX {file_path}: {str(e)}")

        return {
            'text': '\n\n'.join(text_content),
            'metadata': metadata,
            'file_path': file_path
        }


class PPTXLoader(DocumentLoader):
    """Loader for PPTX presentations."""

    def load(self, file_path: str) -> Dict[str, any]:
        """Load PPTX presentation."""
        text_content = []
        metadata = {
            'source': file_path,
            'file_type': 'pptx',
            'num_slides': 0
        }

        try:
            prs = Presentation(file_path)
            metadata['num_slides'] = len(prs.slides)

            # Extract core properties
            core_props = prs.core_properties
            metadata['title'] = core_props.title or ''
            metadata['author'] = core_props.author or ''

            # Extract text from slides
            for slide_num, slide in enumerate(prs.slides):
                slide_text = [f"[Slide {slide_num + 1}]"]

                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_text.append(shape.text)

                if len(slide_text) > 1:  # More than just the slide number
                    text_content.append('\n'.join(slide_text))

        except Exception as e:
            raise RuntimeError(f"Error loading PPTX {file_path}: {str(e)}")

        return {
            'text': '\n\n'.join(text_content),
            'metadata': metadata,
            'file_path': file_path
        }


class XLSXLoader(DocumentLoader):
    """Loader for XLSX spreadsheets."""

    def load(self, file_path: str) -> Dict[str, any]:
        """Load XLSX spreadsheet."""
        text_content = []
        metadata = {
            'source': file_path,
            'file_type': 'xlsx',
            'num_sheets': 0
        }

        try:
            workbook = load_workbook(file_path, data_only=True)
            metadata['num_sheets'] = len(workbook.sheetnames)

            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                sheet_text = [f"[Sheet: {sheet_name}]"]

                for row in sheet.iter_rows(values_only=True):
                    # Filter out completely empty rows
                    row_values = [str(cell) if cell is not None else '' for cell in row]
                    if any(val.strip() for val in row_values):
                        sheet_text.append(' | '.join(row_values))

                if len(sheet_text) > 1:  # More than just the sheet name
                    text_content.append('\n'.join(sheet_text))

        except Exception as e:
            raise RuntimeError(f"Error loading XLSX {file_path}: {str(e)}")

        return {
            'text': '\n\n'.join(text_content),
            'metadata': metadata,
            'file_path': file_path
        }


class TXTLoader(DocumentLoader):
    """Loader for plain text files."""

    def load(self, file_path: str) -> Dict[str, any]:
        """Load plain text file."""
        metadata = {
            'source': file_path,
            'file_type': 'txt'
        }

        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                text = file.read()
        except UnicodeDecodeError:
            # Try with a different encoding if UTF-8 fails
            with open(file_path, 'r', encoding='latin-1') as file:
                text = file.read()
        except Exception as e:
            raise RuntimeError(f"Error loading TXT {file_path}: {str(e)}")

        return {
            'text': text,
            'metadata': metadata,
            'file_path': file_path
        }


class MarkdownLoader(DocumentLoader):
    """Loader for Markdown files."""

    def load(self, file_path: str) -> Dict[str, any]:
        """Load Markdown file."""
        metadata = {
            'source': file_path,
            'file_type': 'markdown'
        }

        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                md_text = file.read()

            # Keep the original markdown for better structure preservation
            # You can also convert to HTML if needed: html = markdown.markdown(md_text)
        except Exception as e:
            raise RuntimeError(f"Error loading Markdown {file_path}: {str(e)}")

        return {
            'text': md_text,
            'metadata': metadata,
            'file_path': file_path
        }


def load_document(file_path: str) -> Dict[str, any]:
    """
    Convenience function to load any supported document.

    Args:
        file_path: Path to the document file

    Returns:
        Dictionary with document content and metadata

    Raises:
        ValueError: If file format is not supported
        RuntimeError: If there's an error loading the document
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    loader = DocumentLoader.get_loader(file_path)
    return loader.load(file_path)
